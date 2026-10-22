"""Excel de órdenes de compra: hoja resumen + una hoja por proveedor."""
from __future__ import annotations

import re
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

AZUL = "1F4E78"
GRIS = "D9E1F2"
BORDE = Border(*[Side(style="thin", color="B4C6E7")] * 4)


def _nombre_hoja(nombre: str, usados: set[str]) -> str:
    limpio = re.sub(r"[\[\]:*?/\\]", " ", str(nombre or "Proveedor")).strip()
    limpio = re.sub(r"\s+", " ", limpio) or "Proveedor"
    base = limpio[:31]
    candidato = base
    i = 2
    while candidato.lower() in usados:
        sufijo = f" ({i})"
        candidato = base[: 31 - len(sufijo)] + sufijo
        i += 1
    usados.add(candidato.lower())
    return candidato


def _estilo_encabezado(ws, fila: int, columnas: int) -> None:
    for c in range(1, columnas + 1):
        celda = ws.cell(row=fila, column=c)
        celda.font = Font(bold=True, color="FFFFFF", size=10)
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center",
                                    wrap_text=True)
        celda.border = BORDE


def _anchos(ws, anchos: list[int]) -> None:
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def exportar_orden_excel(db, orden_id: int, destino: str | Path) -> str:
    """Genera un Excel con una hoja por proveedor. Devuelve la ruta."""
    orden = db.obtener_orden_compra(int(orden_id))
    if not orden:
        raise ValueError(f"No existe la orden #{orden_id}")
    items = orden.get("items") or []
    if not items:
        raise ValueError("La orden no tiene artículos")

    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)

    encabezados = ["Proveedor", "Código", "Descripción", "Cantidad",
                   "Costo unitario", "Subtotal"]
    grupos: "OrderedDict[str, list[dict]]" = OrderedDict()
    for it in items:
        grupos.setdefault(it.get("proveedor_nombre") or "Sin proveedor", []).append(it)

    wb = Workbook()
    resumen = wb.active
    resumen.title = "Resumen"

    resumen["A1"] = f"Orden de compra #{orden_id}"
    resumen["A1"].font = Font(bold=True, size=14, color=AZUL)
    resumen["A2"] = f"Fecha: {str(orden.get('fecha_creacion') or '')[:19]}"
    resumen["A3"] = f"Estado: {orden.get('estado') or ''}"
    resumen["A4"] = f"Observaciones: {orden.get('observaciones') or '-'}"
    resumen.merge_cells("A1:F1")
    resumen.merge_cells("A4:F4")

    fila = 6
    for c, titulo in enumerate(encabezados, start=1):
        resumen.cell(row=fila, column=c, value=titulo)
    _estilo_encabezado(resumen, fila, len(encabezados))

    total_general = 0.0
    fila += 1
    for nombre, lista in grupos.items():
        total_prov = 0.0
        for it in lista:
            valores = [
                nombre,
                it.get("producto_codigo") or "",
                it.get("producto_nombre") or "",
                int(it.get("cantidad") or 0),
                float(it.get("precio_costo") or 0),
                float(it.get("subtotal") or 0),
            ]
            for c, v in enumerate(valores, start=1):
                celda = resumen.cell(row=fila, column=c, value=v)
                celda.border = BORDE
                if c >= 5:
                    celda.number_format = '"$"#,##0.00'
            total_prov += float(it.get("subtotal") or 0)
            fila += 1
        total_general += total_prov
        resumen.cell(row=fila, column=1, value=f"Subtotal {nombre}")
        resumen.cell(row=fila, column=6, value=total_prov).number_format = '"$"#,##0.00'
        for c in range(1, 7):
            resumen.cell(row=fila, column=c).font = Font(bold=True)
            resumen.cell(row=fila, column=c).fill = PatternFill("solid", fgColor=GRIS)
            resumen.cell(row=fila, column=c).border = BORDE
        fila += 1

    resumen.cell(row=fila, column=1, value="TOTAL GENERAL").font = Font(bold=True, size=12)
    celda_total = resumen.cell(row=fila, column=6, value=total_general)
    celda_total.font = Font(bold=True, size=12, color=AZUL)
    celda_total.number_format = '"$"#,##0.00'
    _anchos(resumen, [26, 16, 42, 10, 15, 15])
    resumen.freeze_panes = "A7"

    usados = {"resumen"}
    for nombre, lista in grupos.items():
        ws = wb.create_sheet(_nombre_hoja(nombre, usados))
        ws["A1"] = f"Proveedor: {nombre}"
        ws["A1"].font = Font(bold=True, size=14, color=AZUL)
        ws.merge_cells("A1:F1")
        ws["A2"] = f"Orden #{orden_id} · {str(orden.get('fecha_creacion') or '')[:19]}"
        ws.merge_cells("A2:F2")
        fila = 4
        for c, titulo in enumerate(["Código", "Descripción", "Cantidad",
                                    "Costo unitario", "Subtotal"], start=1):
            ws.cell(row=fila, column=c, value=titulo)
        _estilo_encabezado(ws, fila, 5)
        fila += 1
        total = 0.0
        for it in lista:
            valores = [it.get("producto_codigo") or "", it.get("producto_nombre") or "",
                       int(it.get("cantidad") or 0),
                       float(it.get("precio_costo") or 0),
                       float(it.get("subtotal") or 0)]
            for c, v in enumerate(valores, start=1):
                celda = ws.cell(row=fila, column=c, value=v)
                celda.border = BORDE
                if c >= 4:
                    celda.number_format = '"$"#,##0.00'
            total += float(it.get("subtotal") or 0)
            fila += 1
        ws.cell(row=fila, column=1, value="TOTAL").font = Font(bold=True)
        celda_total = ws.cell(row=fila, column=5, value=total)
        celda_total.font = Font(bold=True, color=AZUL)
        celda_total.number_format = '"$"#,##0.00'
        _anchos(ws, [16, 42, 10, 15, 15])
        ws.freeze_panes = "A5"

    wb.save(destino)
    return str(destino)


def exportar_ordenes_excel(db, ordenes: list[dict], destino: str | Path) -> str:
    """Exporta un listado de órdenes (una fila por orden) a Excel."""
    destino = Path(destino)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Órdenes"
    ws["A1"] = f"Órdenes de compra — {datetime.now():%d/%m/%Y %H:%M}"
    ws["A1"].font = Font(bold=True, size=14, color=AZUL)
    ws.merge_cells("A1:H1")
    encabezados = ["N°", "Fecha", "Estado", "Proveedores", "Ítems", "Unidades",
                   "Total", "Observaciones"]
    for c, t in enumerate(encabezados, start=1):
        ws.cell(row=3, column=c, value=t)
    _estilo_encabezado(ws, 3, len(encabezados))
    fila = 4
    for o in ordenes:
        valores = [o.get("id"), str(o.get("fecha_creacion") or "")[:19],
                   o.get("estado"), o.get("proveedores") or "",
                   o.get("items_count"), o.get("unidades"), o.get("total"),
                   o.get("observaciones") or ""]
        for c, v in enumerate(valores, start=1):
            celda = ws.cell(row=fila, column=c, value=v)
            celda.border = BORDE
            if c == 7:
                celda.number_format = '"$"#,##0.00'
        fila += 1
    _anchos(ws, [8, 20, 13, 32, 8, 10, 15, 40])
    ws.freeze_panes = "A4"
    wb.save(destino)
    return str(destino)
