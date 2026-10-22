"""Reporte de ventas.

- Filtra ventas por rango de fechas (input DD/MM/YYYY).
- Productos más vendidos y ventas por día.
- Exporta a Excel (openpyxl) y a HTML autocontenido con gráficos Apache ECharts.
"""
from __future__ import annotations

import json
import webbrowser
from datetime import datetime
from pathlib import Path
from typing import Tuple


def parse_fecha(fecha_str: str) -> datetime:
    """Convierte 'DD/MM/YYYY' a datetime. Lanza ValueError si es inválida."""
    return datetime.strptime(fecha_str.strip(), "%d/%m/%Y")


def rango_iso(desde: str, hasta: str) -> Tuple[str, str]:
    d = parse_fecha(desde)
    h = parse_fecha(hasta)
    if d > h:
        raise ValueError("La fecha 'desde' no puede ser mayor que la fecha 'hasta'")
    return d.strftime("%Y-%m-%d 00:00:00"), h.strftime("%Y-%m-%d 23:59:59")


def _echarts_js() -> str:
    """Devuelve el JS de ECharts embebido (vendor) o CDN si no está el archivo."""
    base = Path(__file__).resolve().parent.parent / "vendor" / "echarts.min.js"
    try:
        if base.exists():
            return f"<script>{base.read_text(encoding='utf-8', errors='ignore')}</script>"
    except Exception:
        pass
    return '<script src="https://cdn.jsdelivr.net/npm/echarts@5.5.1/dist/echarts.min.js"></script>'


class ReportService:
    @staticmethod
    def resumen(ventas: list) -> tuple[int, float]:
        completadas = [v for v in ventas if v.get("estado") == "completada"]
        return len(completadas), round(sum(v.get("total", 0) for v in completadas), 2)

    @staticmethod
    def exportar_excel(detalle: list, destino: str | Path | None = None) -> str:
        """Exporta el detalle a Excel con encabezado con estilo. Retorna la ruta."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        if not detalle:
            raise ValueError("No hay datos para exportar")
        if destino is None:
            destino = f"reporte_ventas_detallado_{datetime.now():%Y%m%d_%H%M%S}.xlsx"

        wb = Workbook()
        ws = wb.active
        ws.title = "Ventas"
        headers = ["ID Venta", "Fecha", "Producto", "Cantidad", "Precio Unit.", "Subtotal",
                   "Total Venta", "Estado", "Medio pago"]
        ws.append(headers)
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill("solid", fgColor="1F4E79")
        thin = Side(style="thin", color="B0B0B0")
        header_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")
            c.border = header_border
        for row in detalle:
            ws.append([row.get("id"), row.get("fecha"), row.get("nombre_producto"),
                       row.get("cantidad"), row.get("precio_unitario"), row.get("subtotal"),
                       row.get("total"), row.get("estado"), row.get("metodo_pago")])
        widths = [10, 20, 35, 10, 13, 12, 13, 12, 14]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        wb.save(str(destino))
        return str(destino)


    @staticmethod
    def exportar_stock_excel(productos: list, destino: str | Path | None = None) -> str:
        """Exporta el reporte de stock (con estado y prioridad) a Excel."""
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        if not productos:
            raise ValueError("No hay productos para exportar")
        if destino is None:
            destino = f"reporte_stock_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "Stock"
        headers = ["Código", "Descripción", "Categoría", "Proveedor", "Stock",
                   "Mínimo", "Estado", "Prioridad"]
        ws.append(headers)
        thin = Side(style="thin", color="B0B0B0")
        borde = Border(left=thin, right=thin, top=thin, bottom=thin)
        for col in range(1, len(headers) + 1):
            c = ws.cell(row=1, column=col)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E79")
            c.alignment = Alignment(horizontal="center")
            c.border = borde
        estados = {"agotado": "Agotado", "bajo": "Bajo mínimo",
                   "por_agotar": "Por agotar", "ok": "OK"}
        for p in productos:
            ws.append([p.get("codigo"), p.get("descripcion"), p.get("categoria"),
                       p.get("proveedor_nombre") or "Sin proveedor",
                       p.get("stock_actual"), p.get("stock_minimo"),
                       estados.get(p.get("estado_stock"), p.get("estado_stock")),
                       p.get("prioridad_stock")])
        for i, w in enumerate([16, 42, 18, 26, 10, 10, 14, 10], 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
        ws.freeze_panes = "A2"
        wb.save(str(destino))
        return str(destino)

    @staticmethod
    def exportar_html(detalle: list, top: list, por_dia: list, resumen: tuple,
                      destino: str | Path | None = None, abrir: bool = False) -> str:
        """Genera un HTML autocontenido con tablas + gráficos ECharts.

        Args:
            detalle: filas de venta+items (para la tabla de detalle).
            top: productos más vendidos (cantidad_total, monto_total, num_ventas).
            por_dia: ventas por día (dia, cantidad, total).
            resumen: (n_ventas, total).
            destino: ruta de salida (opcional).
            abrir: si True, abre el HTML en el navegador.
        """
        n_ventas, total = resumen
        if destino is None:
            destino = f"reporte_ventas_{datetime.now():%Y%m%d_%H%M%S}.html"
        destino = Path(destino)
        destino.parent.mkdir(parents=True, exist_ok=True)

        top_json = json.dumps(top or [], ensure_ascii=False)
        dia_json = json.dumps([str(d.get("dia")) for d in (por_dia or [])], ensure_ascii=False)
        dia_total_json = json.dumps([float(d.get("total") or 0) for d in (por_dia or [])])
        dia_cant_json = json.dumps([int(d.get("cantidad") or 0) for d in (por_dia or [])])

        filas_top = "".join(
            f"<tr><td>{i+1}</td><td>{str(p.get('descripcion',''))[:40]}</td>"
            f"<td class='num'>{int(p.get('cantidad_total',0))}</td>"
            f"<td class='num'>{int(p.get('num_ventas',0))}</td>"
            f"<td class='num'>${float(p.get('monto_total',0)):,.2f}</td></tr>"
            for i, p in enumerate(top or [])
        )
        filas_det = "".join(
            f"<tr><td>{r.get('id')}</td><td>{str(r.get('fecha',''))[:16]}</td>"
            f"<td>{str(r.get('nombre_producto',''))[:30]}</td>"
            f"<td class='num'>{int(r.get('cantidad',0))}</td>"
            f"<td class='num'>${float(r.get('subtotal',0)):,.2f}</td>"
            f"<td class='num'>${float(r.get('total',0)):,.2f}</td></tr>"
            for r in (detalle or [])
        )

        html = f"""<!doctype html>
<html lang="es"><head><meta charset="utf-8">
<title>Reporte de Ventas</title>
<style>
 body{{font-family:system-ui,Arial,sans-serif;margin:24px;color:#1f2937;background:#f8fafc}}
 h1{{font-size:22px;margin:0 0 4px}} h2{{font-size:16px;margin:24px 0 8px;color:#1f4e79}}
 .resumen{{display:flex;gap:24px;margin:12px 0}}
 .card{{background:#fff;border:1px solid #e5e7eb;border-radius:8px;padding:12px 18px;min-width:160px}}
 .card .n{{font-size:24px;font-weight:700;color:#047857}}
 .card .l{{font-size:12px;color:#6b7280}}
 table{{border-collapse:collapse;width:100%;background:#fff;font-size:13px}}
 th,td{{border:1px solid #e5e7eb;padding:6px 8px;text-align:left}}
 th{{background:#1f4e79;color:#fff}} td.num{{text-align:right;font-variant-numeric:tabular-nums}}
 .grid{{display:grid;grid-template-columns:1fr 1fr;gap:16px}}
 .chart{{height:340px;background:#fff;border:1px solid #e5e7eb;border-radius:8px}}
</style></head>
<body>
<h1>Reporte de Ventas</h1>
<div class="resumen">
  <div class="card"><div class="l">Ventas completadas</div><div class="n">{n_ventas}</div></div>
  <div class="card"><div class="l">Total vendido</div><div class="n">${total:,.2f}</div></div>
  <div class="card"><div class="l">Ticket promedio</div><div class="n">${(total/n_ventas if n_ventas else 0):,.2f}</div></div>
</div>
<h2>Productos más vendidos</h2>
<table><thead><tr><th>#</th><th>Producto</th><th>Cantidad</th><th>Nº ventas</th><th>Monto</th></tr></thead>
<tbody>{filas_top or "<tr><td colspan='5'>Sin datos</td></tr>"}</tbody></table>
<h2>Ventas por día</h2>
<div class="grid">
  <div id="chart-dia" class="chart"></div>
  <div id="chart-top" class="chart"></div>
</div>
<h2>Detalle de ventas</h2>
<table><thead><tr><th>ID</th><th>Fecha</th><th>Producto</th><th>Cant.</th><th>Subtotal</th><th>Total</th></tr></thead>
<tbody>{filas_det or "<tr><td colspan='6'>Sin datos</td></tr>"}</tbody></table>
{_echarts_js()}
<script>
var dias = {dia_json};
var totales = {dia_total_json};
var cantidades = {dia_cant_json};
var top = {top_json};
var topNombres = top.map(function(p){{return String(p.descripcion).slice(0,18)}});
var topCant = top.map(function(p){{return Number(p.cantidad_total)}});
var topMonto = top.map(function(p){{return Number(p.monto_total)}});
var base = {{backgroundColor:'#fff',textStyle:{{fontFamily:'system-ui'}}}};
var c1 = echarts.init(document.getElementById('chart-dia'));
c1.setOption(Object.assign({{}}, base, {{
  title:{{text:'Ventas por día',left:'center'}},
  tooltip:{{trigger:'axis'}},
  xAxis:{{type:'category',data:dias}},
  yAxis:{{type:'value'}},
  series:[
    {{name:'Total $',type:'line',smooth:true,data:totales,areaStyle:{{opacity:.15}}}},
    {{name:'Nº ventas',type:'bar',data:cantidades,yAxisIndex:0}}
  ]
}}));
var c2 = echarts.init(document.getElementById('chart-top'));
c2.setOption(Object.assign({{}}, base, {{
  title:{{text:'Productos más vendidos (cantidad)',left:'center'}},
  tooltip:{{trigger:'axis',axisPointer:{{type:'shadow'}}}},
  grid:{{left:120}},
  xAxis:{{type:'value'}},
  yAxis:{{type:'category',data:topNombres.slice().reverse()}},
  series:[{{type:'bar',data:topCant.slice().reverse(),itemStyle:{{color:'#047857'}}}}]
}}));
window.addEventListener('resize',function(){{c1.resize();c2.resize();}});
</script>
</body></html>"""
        destino.write_text(html, encoding="utf-8")
        if abrir:
            try:
                webbrowser.open(destino.resolve().as_uri())
            except Exception:
                pass
        return str(destino)


report_service = ReportService()
